"""
SYLLABUSIQ - Academic Data Export Service
Supports multi-format exports:
  - JSON (Structured hierarchical syllabus data)
  - CSV (Flat tabular format with UTF-8 BOM)
  - Excel / XLSX (Multi-tab professional workbook with styling)
  - PDF Report (Executive syllabus summary report via ReportLab)
"""

import csv
import io
import json
import logging
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger("syllabusiq.export")


# =========================================================
# HELPER FUNCTIONS
# =========================================================

def _get_course_details(doc: Dict[str, Any], course: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """Retrieve detail object for a given course."""
    details = doc.get("details", {})
    code = (course.get("code") or "").strip()
    title = (course.get("title") or "").strip()

    if code and code in details:
        return details[code]
    if title and title in details:
        return details[title]
    combo = f"{code} | {title}".strip()
    if combo and combo in details:
        return details[combo]

    for k, v in details.items():
        if code and (k == code or k.startswith(code + " ") or k.startswith(code + " |")):
            return v
        if title and title.lower() in k.lower():
            return v
    return None


# =========================================================
# JSON EXPORT
# =========================================================

def export_to_json(doc: Dict[str, Any], options: Optional[Dict[str, Any]] = None) -> bytes:
    """Generate clean, indented JSON export."""
    opts = options or {}
    export_payload = {
        "document_name": doc.get("filename", "syllabus_export.pdf"),
        "college_or_university": doc.get("college", ""),
        "total_pages": doc.get("pages", 0),
        "courses_count": len(doc.get("courses", [])),
        "exported_at": doc.get("uploaded_at", ""),
        "courses": []
    }

    for course in doc.get("courses", []):
        detail = _get_course_details(doc, course) or {}
        c_meta = detail.get("course", {})

        course_item = {
            "code": course.get("code") or c_meta.get("code") or "",
            "title": course.get("title") or c_meta.get("title") or "",
            "category": course.get("category") or c_meta.get("category") or "",
            "credits": course.get("credits") or c_meta.get("credits") or "",
            "semester": course.get("semester") or "",
            "pages": c_meta.get("pages") or [course.get("page", 1)],
            "units": detail.get("units", []),
            "course_outcomes": detail.get("course_outcomes", []),
            "programme_outcomes": detail.get("programme_outcomes", []),
            "co_po_matrix": detail.get("co_po_matrix", {}),
            "books": detail.get("books", [])
        }
        export_payload["courses"].append(course_item)

    return json.dumps(export_payload, ensure_ascii=False, indent=2).encode("utf-8")


# =========================================================
# CSV EXPORT
# =========================================================

def export_to_csv(doc: Dict[str, Any], options: Optional[Dict[str, Any]] = None) -> bytes:
    """Generate structured CSV containing Courses, Units, and Topics."""
    output = io.StringIO()
    writer = csv.writer(output)

    # Header Row
    writer.writerow([
        "Course Code",
        "Course Title",
        "Category",
        "Credits",
        "Page",
        "Unit Number",
        "Unit Title",
        "Unit Hours",
        "Topic ID",
        "Topic / Subtopic Text",
        "Bloom Level",
        "Bloom Source"
    ])

    for course in doc.get("courses", []):
        detail = _get_course_details(doc, course) or {}
        c_meta = detail.get("course", {})
        code = course.get("code") or c_meta.get("code") or ""
        title = course.get("title") or c_meta.get("title") or ""
        category = course.get("category") or c_meta.get("category") or ""
        credits_val = course.get("credits") or c_meta.get("credits") or ""
        page = course.get("page", 1)

        units = detail.get("units", [])
        if not units:
            writer.writerow([code, title, category, credits_val, page, "", "", "", "", "", "", ""])
            continue

        for unit in units:
            u_num = unit.get("number", "")
            u_title = unit.get("title", "")
            u_hours = unit.get("hours", "")
            topics = unit.get("topics", [])

            if not topics:
                writer.writerow([code, title, category, credits_val, page, u_num, u_title, u_hours, "", "", "", ""])
                continue

            for t in topics:
                t_id = t.get("id", "")
                t_text = t.get("text", "")
                t_bloom = t.get("bloom", "")
                t_bsrc = t.get("bloom_source", "")
                writer.writerow([
                    code, title, category, credits_val, page,
                    u_num, u_title, u_hours,
                    t_id, t_text, t_bloom, t_bsrc
                ])

    # Prepend UTF-8 BOM for Excel compatibility
    return ("\ufeff" + output.getvalue()).encode("utf-8")


# =========================================================
# EXCEL (XLSX) EXPORT
# =========================================================

def export_to_excel(doc: Dict[str, Any], options: Optional[Dict[str, Any]] = None) -> bytes:
    """Generate multi-tab styled Excel workbook using openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.warning("openpyxl not installed, falling back to CSV")
        return export_to_csv(doc, options)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="3949C9", end_color="3949C9", fill_type="solid")
    sub_fill = PatternFill(start_color="4B5CD7", end_color="4B5CD7", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0")
    )

    def style_table(ws, headers, fill=header_fill):
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 1. Sheet: Overview
    ws_ov = wb.create_sheet(title="Document Overview")
    ws_ov.views.sheetView[0].showGridLines = True
    ws_ov.append(["SYLLABUSIQ - Academic Syllabus Extraction Report"])
    ws_ov.cell(row=1, column=1).font = Font(size=14, bold=True, color="3949C9")
    ws_ov.append([])
    ws_ov.append(["Property", "Value"])
    ws_ov.append(["Document Name", doc.get("filename", "N/A")])
    ws_ov.append(["College / Institution", doc.get("college", "Universal Syllabus")])
    ws_ov.append(["Total Pages", doc.get("pages", 0)])
    ws_ov.append(["Detected Courses", len(doc.get("courses", []))])
    ws_ov.append(["Status", doc.get("status", "completed")])
    ws_ov.append(["Extraction Engine", "Universal SyllabusIQ Engine (LLaMA 3.2 Vision + PyMuPDF)"])

    # 2. Sheet: Courses Summary
    ws_courses = wb.create_sheet(title="Courses")
    ws_courses.views.sheetView[0].showGridLines = True
    style_table(ws_courses, ["Course Code", "Course Title", "Category", "Credits", "Semester", "Start Page", "Units Count", "Outcomes Count"])

    for c in doc.get("courses", []):
        det = _get_course_details(doc, c) or {}
        c_meta = det.get("course", {})
        ws_courses.append([
            c.get("code") or c_meta.get("code") or "",
            c.get("title") or c_meta.get("title") or "",
            c.get("category") or c_meta.get("category") or "",
            c.get("credits") or c_meta.get("credits") or "",
            c.get("semester") or "",
            c.get("page", 1),
            len(det.get("units", [])),
            len(det.get("course_outcomes", []))
        ])

    # 3. Sheet: Units & Topics
    ws_topics = wb.create_sheet(title="Units & Topics")
    ws_topics.views.sheetView[0].showGridLines = True
    style_table(ws_topics, ["Course Code", "Course Title", "Unit #", "Unit Title", "Hours", "Topic ID", "Topic / Subtopic Description", "Bloom Level", "Inference Type"])

    for c in doc.get("courses", []):
        det = _get_course_details(doc, c) or {}
        code = c.get("code") or ""
        title = c.get("title") or ""
        for u in det.get("units", []):
            u_num = u.get("number", "")
            u_title = u.get("title", "")
            u_hrs = u.get("hours", "")
            topics = u.get("topics", [])
            if not topics:
                ws_topics.append([code, title, u_num, u_title, u_hrs, "", "", "", ""])
            for t in topics:
                ws_topics.append([
                    code, title, u_num, u_title, u_hrs,
                    t.get("id", ""),
                    t.get("text", ""),
                    t.get("bloom", ""),
                    t.get("bloom_source", "")
                ])

    # 4. Sheet: Course Outcomes
    ws_cos = wb.create_sheet(title="Course Outcomes")
    ws_cos.views.sheetView[0].showGridLines = True
    style_table(ws_cos, ["Course Code", "Course Title", "CO ID", "Outcome Description", "Bloom Level", "Bloom Source"])

    for c in doc.get("courses", []):
        det = _get_course_details(doc, c) or {}
        code = c.get("code") or ""
        title = c.get("title") or ""
        for co in det.get("course_outcomes", []):
            ws_cos.append([
                code, title,
                co.get("id", ""),
                co.get("text", ""),
                co.get("bloom", ""),
                co.get("bloom_source", "")
            ])

    # 5. Sheet: Books & References
    ws_books = wb.create_sheet(title="Books & References")
    ws_books.views.sheetView[0].showGridLines = True
    style_table(ws_books, ["Course Code", "Course Title", "Type", "Book Title", "Author", "Publisher", "Edition / Year"])

    for c in doc.get("courses", []):
        det = _get_course_details(doc, c) or {}
        code = c.get("code") or ""
        title = c.get("title") or ""
        for b in det.get("books", []):
            ws_books.append([
                code, title,
                b.get("kind", "text").upper(),
                b.get("title", ""),
                b.get("author", ""),
                b.get("publisher", ""),
                f"{b.get('edition', '')} {b.get('year', '')}".strip()
            ])

    # Auto-fit column widths
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                max_len = max(max_len, len(val))
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =========================================================
# PDF REPORT EXPORT
# =========================================================

def export_to_pdf(doc: Dict[str, Any], options: Optional[Dict[str, Any]] = None) -> bytes:
    """Generate executive PDF summary report using ReportLab."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
    except ImportError:
        logger.warning("reportlab not installed, falling back to JSON")
        return export_to_json(doc, options)

    buffer = io.BytesIO()
    doc_pdf = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#3949C9'),
        spaceAfter=12
    )
    heading2_style = ParagraphStyle(
        'DocHeading2',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=14,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155')
    )

    story = []

    # Title
    story.append(Paragraph("SYLLABUSIQ — Academic Extraction Report", title_style))
    story.append(Paragraph(f"<b>Document:</b> {doc.get('filename', 'Syllabus')} | <b>Total Pages:</b> {doc.get('pages', 0)} | <b>Courses:</b> {len(doc.get('courses', []))}", body_style))
    story.append(Spacer(1, 14))

    # Courses Summary Table
    courses_data = [["Code", "Course Title", "Units", "Outcomes", "Start Page"]]
    for c in doc.get("courses", []):
        det = _get_course_details(doc, c) or {}
        courses_data.append([
            c.get("code", "—"),
            Paragraph(c.get("title", "Untitled"), body_style),
            str(len(det.get("units", []))),
            str(len(det.get("course_outcomes", []))),
            str(c.get("page", "1"))
        ])

    t = Table(courses_data, colWidths=[80, 240, 60, 70, 70])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3949C9')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    story.append(t)
    story.append(Spacer(1, 16))

    # Detailed Course Breakdowns
    for c in doc.get("courses", []):
        det = _get_course_details(doc, c)
        if not det:
            continue

        c_title = f"{c.get('code', '')} {c.get('title', '')}".strip()
        story.append(Paragraph(f"Course: {c_title}", heading2_style))

        units = det.get("units", [])
        if units:
            unit_data = [["Unit", "Title & Topics", "Hours"]]
            for u in units:
                topics_text = "<br/>".join([f"• {t.get('text', '')}" for t in u.get("topics", [])])
                content = f"<b>{u.get('title', '')}</b>" + (f"<br/>{topics_text}" if topics_text else "")
                unit_data.append([
                    f"Unit {u.get('number', '')}",
                    Paragraph(content, body_style),
                    f"{u.get('hours', '—')} hrs"
                ])

            ut = Table(unit_data, colWidths=[60, 400, 60])
            ut.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]))
            story.append(ut)
            story.append(Spacer(1, 10))

    doc_pdf.build(story)
    return buffer.getvalue()


# =========================================================
# UNIFIED DISPATCHER
# =========================================================

def generate_export(
    doc: Dict[str, Any],
    format_name: str = "json",
    options: Optional[Dict[str, Any]] = None
) -> Tuple[bytes, str, str]:
    """
    Export syllabus data in the requested format.
    Returns: (bytes_content, media_type, file_extension)
    """
    fmt = (format_name or "json").lower().strip()

    if fmt in ["xlsx", "excel"]:
        content = export_to_excel(doc, options)
        return content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"

    elif fmt in ["csv"]:
        content = export_to_csv(doc, options)
        return content, "text/csv; charset=utf-8", "csv"

    elif fmt in ["pdf"]:
        content = export_to_pdf(doc, options)
        return content, "application/pdf", "pdf"

    else:
        content = export_to_json(doc, options)
        return content, "application/json", "json"
